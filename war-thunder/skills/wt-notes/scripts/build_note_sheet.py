#!/usr/bin/env python3
"""
build_note_sheet.py — Reconstruit à l'identique la feuille "note" du classeur
"Notes avions WT" de Florian (système de notation d'avions War Thunder).

Ce script ne construit QUE la feuille "note" (grille de notation). Il ne touche
à aucune autre feuille du classeur (Feuil1, Avion WT, Avion WT (2), Feuil2,
Feuil3, Claude) — si elles existent déjà dans le classeur cible, elles restent
intactes ; si la feuille "note" existe déjà, elle est supprimée et recréée.

Usage :
    python build_note_sheet.py <classeur.xlsx> [classeur_sortie.xlsx]

Si aucun fichier de sortie n'est précisé, le classeur d'entrée est écrasé.

⚠️ IMPORTANT — ce script ne fait qu'écrire les formules, il ne les calcule
jamais (openpyxl n'exécute aucune formule). Après exécution, il FAUT
recalculer le classeur via LibreOffice headless avant de le livrer, sinon les
cellules de formule (◆ Note, zone de calcul masquée) apparaîtront vides à
l'ouverture. Voir la section "Pièges Excel/LibreOffice" du SKILL.md associé
pour la marche à suivre complète (recalcul, ré-injection de fullCalcOnLoad,
réparation du filtre de Feuil2 si le classeur cible le contient).
"""

import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

# ---------------------------------------------------------------------------
# Constantes de mise en page
# ---------------------------------------------------------------------------
R_BANNER, R_NOTE, R_SUB, R_COEF, R_DATA0, N_DATA = 1, 2, 3, 4, 5, 40

# (couleur bannière catégorie, couleur note, couleur sous-note) par catégorie
PAL = {
    'avio': ("375623", "548235", "C6E0B4"),  # vert
    'mani': ("1F4E78", "2E75B6", "BDD7EE"),  # bleu
    'mote': ("843C0C", "C55A11", "F8CBAD"),  # orange
    'resi': ("3B3838", "757070", "D0CECE"),  # gris
    'offe': ("53307E", "8064A2", "D9C9EA"),  # violet
    'susp': ("806000", "BF8F00", "FFE699"),  # or
    'defe': ("7B2E2E", "A94D4D", "E6B8B7"),  # rouge
    'role': ("2D2D5F", "44457F", "C5C5E0"),  # indigo
}
NEUTRAL, COEF_FILL, WHITE, BLACK = "404040", "F2F2F2", "FFFFFF", "000000"
THIN = Side(style="thin", color="D9D9D9")
BORDER_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def _set_cell(ws, coord, value=None, *, bold=False, size=10, color=BLACK, bg=None,
              rot=0, align=None, valign="center", wrap=True, border=True, italic=False):
    c = ws[coord] if isinstance(coord, str) else coord
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", size=size, bold=bold, color=color, italic=italic)
    if bg:
        c.fill = _fill(bg)
    if align is None:
        align = "center"
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap, text_rotation=rot)
    if border:
        c.border = BORDER_THIN
    return c


def _helper_formula(src_ref):
    """Formule de la colonne masquée : extrait objectif+subjectif d'une cellule
    saisie au format texte '5+2' / '4-1' / '6' / (vide). Renvoie 0 (jamais "")
    pour une case vide — condition nécessaire pour que la moyenne pondérée en
    aval reste une multiplication purement arithmétique (voir avertissement
    ci-dessous)."""
    return (f'=IF({src_ref}="",0,'
            f'IFERROR(VALUE(LEFT({src_ref},FIND("+",{src_ref})-1))+VALUE(MID({src_ref},FIND("+",{src_ref})+1,10)),'
            f'IFERROR(VALUE(LEFT({src_ref},FIND("-",{src_ref})-1))-VALUE(MID({src_ref},FIND("-",{src_ref})+1,10)),'
            f'VALUE({src_ref}))))')


def build_note_sheet(wb):
    """Construit (ou reconstruit) la feuille "note" dans le classeur openpyxl
    `wb` fourni, et la retourne. Ne modifie aucune autre feuille."""

    if 'note' in wb.sheetnames:
        del wb['note']
    ws = wb.create_sheet('note')

    # -----------------------------------------------------------------
    # Disposition des colonnes
    # -----------------------------------------------------------------
    fixed = ["Avion", "Nation", "BR", "Rang", "Type"]
    col = 1
    FIX = {}
    for name in fixed:
        FIX[name] = col
        col += 1

    blocks = []

    def simple_block(key, cat_label, leaves):
        nonlocal col
        start = col
        score_col = col
        col += 1
        leaf_cols = []
        for (lab, coef) in leaves:
            leaf_cols.append((col, lab, coef))
            col += 1
        end = col - 1
        blocks.append(dict(kind='simple', key=key, cat=cat_label, score_col=score_col,
                            leaves=leaf_cols, start=start, end=end))

    # NOTE : ⁽ᵈ⁾ = valeur dupliquée depuis Avionique. ⁽ᵖ⁾ = critère "proposition",
    # pas encore validé par Florian — voir SKILL.md, section "Décisions ouvertes".
    simple_block('avio', "Avionique",
        [("Volets", 3), ("Pas d'Hélice", 3), ("MEC", 4), ("Trim", 1),
         ("Radiateur et Refroidissement", 3)])
    simple_block('mani', "Maniabilité",
        [("Virage à Plat", 4), ("Roulis", 4), ("Efficacité des Volets", 3),
         ("Conservation d'Énergie en virage", 3), ("Performance verticale ⁽ᵖ⁾", 3)])
    simple_block('mote', "Moteur et Énergie",
        [("Puissance du Moteur (spéc.)", 5), ("Taux de Montée", 5),
         ("Vitesse Max. Palier", 4), ("Vitesse Max. Piqué", 4),
         ("Pas d'Hélice ⁽ᵈ⁾", 2), ("MEC ⁽ᵈ⁾", 2), ("Radiateur & Refr. ⁽ᵈ⁾", 2),
         ("Altitude critique ⁽ᵖ⁾", 3)])
    simple_block('resi', "Résilience",
        [("Blindage de l'Équipage", 4), ("Protection de la cellule", 3),
         ("Protection modules", 4), ("Résistance du Bloc Moteur", 3),
         ("Risque Incendie ⁽ᵖ⁾", 2)])

    off_start = col
    off_score = col; col += 1
    pf_score = col; col += 1
    armes = []
    for lab in ["Armes 1", "Armes 2", "Armes 3 ou Pods"]:
        armes.append((col, lab, 1)); col += 1
    div_col = col; col += 1
    prec_col = col; col += 1
    qm_score = col; col += 1
    muns = []
    for lab in ["Munitions 1", "Munitions 2", "Munitions 3 ou Pods"]:
        muns.append((col, lab, 1)); col += 1
    pos_col = col; col += 1
    off_end = col - 1
    blocks.append(dict(kind='offe', cat="Armement Offensif",
                        off_score=off_score, pf_score=pf_score, armes=armes,
                        div=div_col, prec=prec_col, qm_score=qm_score, muns=muns,
                        pos=pos_col, start=off_start, end=off_end))

    simple_block('susp', "Armement Suspendu",
        [("Charge Utile Maximale", 1), ("Flexibilité de l'Armement", 1)])
    simple_block('defe', "Armement Défensif",
        [("Puissance Armes Défensives", 1), ("Angle de Tir", 2)])

    role_start = col
    roles = ["Chasseur Dogfight", "Boom & Zoom", "Intercepteur",
             "Bombardier en piqué", "Attaqueur au sol", "Bombardier tactique",
             "Bombardier Stratégique"]
    role_cols = []
    for r in roles:
        role_cols.append((col, r)); col += 1
    role_end = col - 1
    LAST_COL = col - 1

    # Une colonne masquée par note manuelle, pour parser "obj+subj"
    leaf_columns = []
    for b in blocks:
        if b['kind'] == 'simple':
            for (cidx, lab, coef) in b['leaves']:
                leaf_columns.append(cidx)
        elif b['kind'] == 'offe':
            for (cidx, lab, coef) in b['armes']:
                leaf_columns.append(cidx)
            leaf_columns.append(b['div'])
            leaf_columns.append(b['prec'])
            for (cidx, lab, coef) in b['muns']:
                leaf_columns.append(cidx)
            leaf_columns.append(b['pos'])

    HELPER_START = LAST_COL + 3
    helper_of = {}
    for i, vcol in enumerate(leaf_columns):
        helper_of[vcol] = HELPER_START + i
    HELPER_END = HELPER_START + len(leaf_columns) - 1

    # -----------------------------------------------------------------
    # En-têtes colonnes fixes
    # -----------------------------------------------------------------
    for name, cidx in FIX.items():
        cl = L(cidx)
        ws.merge_cells(f"{cl}{R_BANNER}:{cl}{R_COEF}")
        _set_cell(ws, f"{cl}{R_BANNER}", name, bold=True, size=11, color=WHITE, bg=NEUTRAL)

    # -----------------------------------------------------------------
    # En-têtes des catégories (bannière + notes + coefficients)
    # -----------------------------------------------------------------
    for b in blocks:
        if b['kind'] == 'simple':
            dark, midc, light = PAL[b['key']]
            ws.merge_cells(f"{L(b['start'])}{R_BANNER}:{L(b['end'])}{R_BANNER}")
            _set_cell(ws, f"{L(b['start'])}{R_BANNER}", b['cat'], bold=True, size=11, color=WHITE, bg=dark)
            sc = L(b['score_col'])
            ws.merge_cells(f"{sc}{R_NOTE}:{sc}{R_SUB}")
            _set_cell(ws, f"{sc}{R_NOTE}", "◆ Note", bold=True, size=9, color=WHITE, bg=dark, rot=90)
            _set_cell(ws, f"{sc}{R_COEF}", "", bg=COEF_FILL)
            for (cidx, lab, coef) in b['leaves']:
                cl = L(cidx)
                ws.merge_cells(f"{cl}{R_NOTE}:{cl}{R_SUB}")
                _set_cell(ws, f"{cl}{R_NOTE}", lab, size=9, color=WHITE, bg=midc, rot=90)
                _set_cell(ws, f"{cl}{R_COEF}", coef, size=9, bg=COEF_FILL, italic=True)
        elif b['kind'] == 'offe':
            dark, midc, light = PAL['offe']
            ws.merge_cells(f"{L(b['start'])}{R_BANNER}:{L(b['end'])}{R_BANNER}")
            _set_cell(ws, f"{L(b['start'])}{R_BANNER}", b['cat'], bold=True, size=11, color=WHITE, bg=dark)
            sc = L(b['off_score'])
            ws.merge_cells(f"{sc}{R_NOTE}:{sc}{R_SUB}")
            _set_cell(ws, f"{sc}{R_NOTE}", "◆ Note", bold=True, size=9, color=WHITE, bg=dark, rot=90)
            _set_cell(ws, f"{sc}{R_COEF}", "", bg=COEF_FILL)
            pf = b['pf_score']; arme_last = b['armes'][-1][0]
            ws.merge_cells(f"{L(pf)}{R_NOTE}:{L(arme_last)}{R_NOTE}")
            _set_cell(ws, f"{L(pf)}{R_NOTE}", "Puissance de Feu", bold=True, size=9, color=WHITE, bg=midc, rot=0)
            _set_cell(ws, f"{L(pf)}{R_SUB}", "◆ Note", bold=True, size=8, color=WHITE, bg=dark, rot=90)
            _set_cell(ws, f"{L(pf)}{R_COEF}", 5, size=9, bg=COEF_FILL, italic=True)
            for (cidx, lab, coef) in b['armes']:
                _set_cell(ws, f"{L(cidx)}{R_SUB}", lab, size=8, color=BLACK, bg=light, rot=90)
                _set_cell(ws, f"{L(cidx)}{R_COEF}", coef, size=9, bg=COEF_FILL, italic=True)
            for cidx, lab, coef in [(b['div'], "Diversité de l'Armement", 2), (b['prec'], "Précision", 4)]:
                cl = L(cidx)
                ws.merge_cells(f"{cl}{R_NOTE}:{cl}{R_SUB}")
                _set_cell(ws, f"{cl}{R_NOTE}", lab, size=9, color=WHITE, bg=midc, rot=90)
                _set_cell(ws, f"{cl}{R_COEF}", coef, size=9, bg=COEF_FILL, italic=True)
            qm = b['qm_score']; mun_last = b['muns'][-1][0]
            ws.merge_cells(f"{L(qm)}{R_NOTE}:{L(mun_last)}{R_NOTE}")
            _set_cell(ws, f"{L(qm)}{R_NOTE}", "Quantité de munitions", bold=True, size=9, color=WHITE, bg=midc, rot=0)
            _set_cell(ws, f"{L(qm)}{R_SUB}", "◆ Note", bold=True, size=8, color=WHITE, bg=dark, rot=90)
            _set_cell(ws, f"{L(qm)}{R_COEF}", 4, size=9, bg=COEF_FILL, italic=True)
            for (cidx, lab, coef) in b['muns']:
                _set_cell(ws, f"{L(cidx)}{R_SUB}", lab, size=8, color=BLACK, bg=light, rot=90)
                _set_cell(ws, f"{L(cidx)}{R_COEF}", coef, size=9, bg=COEF_FILL, italic=True)
            cl = L(b['pos'])
            ws.merge_cells(f"{cl}{R_NOTE}:{cl}{R_SUB}")
            _set_cell(ws, f"{cl}{R_NOTE}", "Position des armes ⁽ᵖ⁾", size=9, color=WHITE, bg=midc, rot=90)
            _set_cell(ws, f"{cl}{R_COEF}", 2, size=9, bg=COEF_FILL, italic=True)

    dark, midc, light = PAL['role']
    ws.merge_cells(f"{L(role_start)}{R_BANNER}:{L(role_end)}{R_BANNER}")
    _set_cell(ws, f"{L(role_start)}{R_BANNER}", "NOTES DE RÔLE (à définir)", bold=True, size=11, color=WHITE, bg=dark)
    for (cidx, lab) in role_cols:
        cl = L(cidx)
        ws.merge_cells(f"{cl}{R_NOTE}:{cl}{R_SUB}")
        _set_cell(ws, f"{cl}{R_NOTE}", lab, bold=True, size=9, color=WHITE, bg=midc, rot=90)
        _set_cell(ws, f"{cl}{R_COEF}", "±4", size=8, color="C00000", bg=COEF_FILL, italic=True)

    ws.cell(1, HELPER_START, "ZONE DE CALCUL (masquée) — extraction obj+subj de chaque note").font = \
        Font(size=9, italic=True, color="808080")
    for hc in range(HELPER_START, HELPER_END + 1):
        ws.column_dimensions[L(hc)].width = 5
        ws.column_dimensions[L(hc)].outlineLevel = 1
        ws.column_dimensions[L(hc)].hidden = True

    # -----------------------------------------------------------------
    # Lignes de données (40 lignes vides prêtes à l'emploi) + formules
    # -----------------------------------------------------------------
    def rng(c1, c2, row):
        return f"{L(c1)}{row}:{L(c2)}{row}"

    for i in range(N_DATA):
        row = R_DATA0 + i
        for name, cidx in FIX.items():
            _set_cell(ws, f"{L(cidx)}{row}", None, align="center")

        for vcol in leaf_columns:
            hc = helper_of[vcol]
            src = f"{L(vcol)}{row}"
            _set_cell(ws, f"{L(hc)}{row}", _helper_formula(src), size=9, border=True)

        for b in blocks:
            if b['kind'] == 'simple':
                leaf_ids = [lc[0] for lc in b['leaves']]
                c1, c2 = leaf_ids[0], leaf_ids[-1]
                h1, h2 = helper_of[c1], helper_of[c2]
                vis = rng(c1, c2, row)
                hel = rng(h1, h2, row)
                coefr = f"{L(c1)}${R_COEF}:{L(c2)}${R_COEF}"
                # ⚠️ Arithmétique pure uniquement (coef * bool * valeur), AUCUN
                # appel de fonction (N(), IF() sur une plage) à l'intérieur du
                # SUMPRODUCT : ça casse silencieusement dans Excel (intersection
                # implicite, seule la 1ère cellule de la plage est utilisée).
                f = f'=IFERROR(SUMPRODUCT({coefr}*({vis}<>"")*{hel})/SUMPRODUCT({coefr}*({vis}<>"")),"")'
                sc = L(b['score_col'])
                c = _set_cell(ws, f"{sc}{row}", f, bold=True, bg="F7F7F7")
                c.number_format = "0.0"
                for (cidx, lab, coef) in b['leaves']:
                    _set_cell(ws, f"{L(cidx)}{row}", None, align="center")
                    ws.cell(row, cidx).number_format = '@'  # Texte : évite l'auto-conversion en date ("4-1" -> 4 janvier)
            elif b['kind'] == 'offe':
                a1, a3 = b['armes'][0][0], b['armes'][-1][0]
                ha1, ha3 = helper_of[a1], helper_of[a3]
                vis = rng(a1, a3, row); hel = rng(ha1, ha3, row)
                coefr = f"{L(a1)}${R_COEF}:{L(a3)}${R_COEF}"
                fpf = f'=IFERROR(SUMPRODUCT({coefr}*({vis}<>"")*{hel})/SUMPRODUCT({coefr}*({vis}<>"")),"")'
                c = _set_cell(ws, f"{L(b['pf_score'])}{row}", fpf, bold=True, bg="F7F7F7")
                c.number_format = "0.0"
                for (cidx, lab, coef) in b['armes']:
                    _set_cell(ws, f"{L(cidx)}{row}", None)
                    ws.cell(row, cidx).number_format = '@'

                m1, m3 = b['muns'][0][0], b['muns'][-1][0]
                hm1, hm3 = helper_of[m1], helper_of[m3]
                vis = rng(m1, m3, row); hel = rng(hm1, hm3, row)
                coefr = f"{L(m1)}${R_COEF}:{L(m3)}${R_COEF}"
                fqm = f'=IFERROR(SUMPRODUCT({coefr}*({vis}<>"")*{hel})/SUMPRODUCT({coefr}*({vis}<>"")),"")'
                c = _set_cell(ws, f"{L(b['qm_score'])}{row}", fqm, bold=True, bg="F7F7F7")
                c.number_format = "0.0"
                for (cidx, lab, coef) in b['muns']:
                    _set_cell(ws, f"{L(cidx)}{row}", None)
                    ws.cell(row, cidx).number_format = '@'

                _set_cell(ws, f"{L(b['div'])}{row}", None); ws.cell(row, b['div']).number_format = '@'
                _set_cell(ws, f"{L(b['prec'])}{row}", None); ws.cell(row, b['prec']).number_format = '@'
                _set_cell(ws, f"{L(b['pos'])}{row}", None); ws.cell(row, b['pos']).number_format = '@'

                PF = L(b['pf_score']); DV = L(b['div']); PR = L(b['prec']); QM = L(b['qm_score']); PO = L(b['pos'])
                hDV = L(helper_of[b['div']]); hPR = L(helper_of[b['prec']]); hPO = L(helper_of[b['pos']])
                cf = R_COEF
                num = (f'IF({PF}{row}="",0,{PF}{row}*${PF}${cf})'
                       f'+IF({DV}{row}="",0,{hDV}{row}*${DV}${cf})'
                       f'+IF({PR}{row}="",0,{hPR}{row}*${PR}${cf})'
                       f'+IF({QM}{row}="",0,{QM}{row}*${QM}${cf})'
                       f'+IF({PO}{row}="",0,{hPO}{row}*${PO}${cf})')
                den = (f'IF({PF}{row}="",0,${PF}${cf})'
                       f'+IF({DV}{row}="",0,${DV}${cf})'
                       f'+IF({PR}{row}="",0,${PR}${cf})'
                       f'+IF({QM}{row}="",0,${QM}${cf})'
                       f'+IF({PO}{row}="",0,${PO}${cf})')
                foff = f'=IFERROR(({num})/({den}),"")'
                c = _set_cell(ws, f"{L(b['off_score'])}{row}", foff, bold=True, bg="F7F7F7")
                c.number_format = "0.0"

        for (cidx, lab) in role_cols:
            _set_cell(ws, f"{L(cidx)}{row}", None, align="center")

    # -----------------------------------------------------------------
    # Dimensions, gel des volets, mise en page
    # -----------------------------------------------------------------
    ws.column_dimensions[L(FIX['Avion'])].width = 20
    ws.column_dimensions[L(FIX['Nation'])].width = 12
    ws.column_dimensions[L(FIX['BR'])].width = 6
    ws.column_dimensions[L(FIX['Rang'])].width = 7
    ws.column_dimensions[L(FIX['Type'])].width = 11
    for cidx in range(FIX['Type'] + 1, LAST_COL + 1):
        ws.column_dimensions[L(cidx)].width = 4.6

    ws.row_dimensions[R_BANNER].height = 22
    ws.row_dimensions[R_NOTE].height = 128
    ws.row_dimensions[R_SUB].height = 66
    ws.row_dimensions[R_COEF].height = 16
    for i in range(N_DATA):
        ws.row_dimensions[R_DATA0 + i].height = 18

    ws.freeze_panes = f"{L(FIX['Type'] + 1)}{R_DATA0}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'

    return ws


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python build_note_sheet.py <classeur.xlsx> [classeur_sortie.xlsx]")
        sys.exit(1)

    in_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else in_path

    wb = openpyxl.load_workbook(in_path)
    build_note_sheet(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(f"Feuille 'note' (re)construite dans {out_path}.")
    print("⚠️  Pense à recalculer via LibreOffice avant de livrer le fichier "
          "(voir SKILL.md, section Pièges Excel/LibreOffice).")
